import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from fastapi import HTTPException, UploadFile
from sqlalchemy import Boolean, Column, Date, DateTime, Integer, JSON, MetaData, Numeric, String, Table, inspect, insert, select, update, delete, func
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.dataset import Dataset, DatasetRevision, IngestionRun, UploadSession
from app.models.workspace import Workspace


SQL_TYPES = {"text": String, "integer": Integer, "decimal": Numeric, "boolean": Boolean, "date": Date, "datetime": DateTime}


def workspace_for_user(db: Session, workspace_id: int, user_id: int) -> Workspace:
    workspace = db.query(Workspace).filter(Workspace.id == workspace_id, Workspace.owner_id == user_id).first()
    if workspace is None:
        raise HTTPException(404, "Workspace not found")
    return workspace


def owned_dataset(db: Session, dataset_id: int, user_id: int, for_update: bool = False) -> Dataset:
    query = db.query(Dataset).join(Workspace).filter(Dataset.id == dataset_id, Workspace.owner_id == user_id)
    if for_update:
        query = query.with_for_update()
    dataset = query.first()
    if dataset is None:
        raise HTTPException(404, "Dataset not found")
    return dataset


def validate_schema(schema: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not schema or len(schema) > get_settings().upload_max_columns:
        raise HTTPException(422, "Dataset must have between 1 and the configured maximum number of columns")
    names: set[str] = set()
    cleaned = []
    for item in schema:
        name = str(item.get("name", "")).strip()
        kind = item.get("type", "text")
        if not name or len(name) > 128 or name.casefold() in names or kind not in SQL_TYPES:
            raise HTTPException(422, "Column names must be unique and column types must be supported")
        names.add(name.casefold())
        cleaned.append({"name": name, "type": kind, "nullable": bool(item.get("nullable", True))})
    return cleaned


def managed_table(dataset: Dataset, create: bool = False) -> Table:
    if not dataset.storage_table_key or not re.fullmatch(r"md_[0-9a-f]{32}", dataset.storage_table_key):
        raise HTTPException(409, "Dataset storage is unavailable")
    metadata = MetaData()
    cols = [Column("row_id", String(32), primary_key=True), Column("workspace_id", Integer, nullable=False, index=True)]
    for spec in dataset.schema_json:
        key = spec.get("key")
        if not isinstance(key, str) or not re.fullmatch(r"c_[0-9a-f]{32}", key):
            raise HTTPException(409, "Dataset schema is invalid")
        cols.append(Column(key, SQL_TYPES[spec["type"]](), nullable=bool(spec.get("nullable", True))))
    return Table(dataset.storage_table_key, metadata, *cols)


def make_schema(schema: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{**item, "key": f"c_{uuid.uuid4().hex}"} for item in validate_schema(schema)]


def create_storage(db: Session, dataset: Dataset) -> None:
    table = managed_table(dataset)
    table.create(db.connection(), checkfirst=True)


def coerce_value(value: Any, kind: str, nullable: bool) -> Any:
    if value is None or value == "":
        if nullable:
            return None
        raise ValueError("A required value is missing")
    if kind == "text":
        return str(value)
    if kind == "integer":
        if isinstance(value, bool): raise ValueError("Expected an integer")
        if isinstance(value, float) and value.is_integer(): return int(value)
        number = int(value)
        if str(number) != str(value).strip() and not isinstance(value, int): raise ValueError("Expected an integer")
        return number
    if kind == "decimal":
        number = Decimal(str(value))
        if not number.is_finite(): raise ValueError("Expected a finite decimal")
        return number
    if kind == "boolean":
        if isinstance(value, bool): return value
        if str(value).casefold() in {"true", "1", "yes"}: return True
        if str(value).casefold() in {"false", "0", "no"}: return False
        raise ValueError("Expected a boolean")
    if kind == "date":
        if isinstance(value, datetime): return value.date()
        return value if isinstance(value, date) else date.fromisoformat(str(value))
    if kind == "datetime":
        if isinstance(value, datetime): return value
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    raise ValueError("Unsupported column type")


def validate_rows(dataset: Dataset, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output = []
    errors = []
    specs = {column["name"]: column for column in dataset.schema_json}
    for row_number, row in enumerate(rows, start=1):
        if set(row) - set(specs):
            errors.append({"row": row_number, "column": "", "message": "Unexpected columns"})
        mapped = {"row_id": uuid.uuid4().hex}
        for name, spec in specs.items():
            try:
                mapped[spec["key"]] = coerce_value(row.get(name), spec["type"], spec["nullable"])
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                errors.append({"row": row_number, "column": name, "message": "Value does not match the declared type or nullability"})
        output.append(mapped)
    if errors:
        raise HTTPException(422, detail={"message": "Some rows failed validation", "errors": errors[:100]})
    return output


def record_revision(db: Session, dataset: Dataset) -> None:
    last = db.query(func.max(DatasetRevision.version)).filter(DatasetRevision.dataset_id == dataset.id).scalar() or 0
    db.add(DatasetRevision(dataset_id=dataset.id, version=last + 1, schema_snapshot=dataset.schema_json, row_count=dataset.row_count))


def safe_filename(filename: str | None) -> str:
    value = (filename or "upload").replace("\\", "/").split("/")[-1]
    value = "".join(ch for ch in value if ch.isprintable() and ch not in "\r\n\x00")
    return value[:255] or "upload"


def staging_path(key: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", key):
        raise HTTPException(404, "Upload not found")
    root = Path(get_settings().upload_staging_dir).resolve()
    path = (root / f"{key}.bin").resolve()
    if path.parent != root:
        raise HTTPException(404, "Upload not found")
    return path


async def save_upload(db: Session, workspace_id: int, user_id: int, upload: UploadFile) -> UploadSession:
    workspace_for_user(db, workspace_id, user_id)
    filename = safe_filename(upload.filename)
    suffix = Path(filename).suffix.lower()
    file_type = {".csv": "csv", ".xlsx": "xlsx", ".xls": "xls", ".json": "json"}.get(suffix)
    if not file_type:
        raise HTTPException(415, "Supported files are CSV, XLSX, XLS, and JSON")
    settings = get_settings()
    key = uuid.uuid4().hex
    path = staging_path(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    if os.name != "nt":
        path.parent.chmod(0o700)
    size = 0
    digest = hashlib.sha256()
    try:
        with path.open("xb") as target:
            while chunk := await upload.read(1024 * 1024):
                size += len(chunk)
                if size > settings.upload_max_bytes:
                    raise HTTPException(413, "File exceeds the configured upload size limit")
                digest.update(chunk)
                target.write(chunk)
        if not size:
            raise HTTPException(422, "The uploaded file is empty")
        sheets = []
        if file_type == "xlsx":
            with zipfile.ZipFile(path) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(item.file_size for item in infos) > settings.upload_max_bytes * 20:
                    raise HTTPException(413, "Excel file expands beyond the allowed limit")
                if any(item.filename.lower().endswith("vba_project.bin") for item in infos):
                    raise HTTPException(415, "Macro-enabled workbooks are not supported")
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheets = book.sheetnames
            book.close()
        elif file_type == "xls":
            book = xlrd.open_workbook(path, on_demand=True)
            sheets = book.sheet_names()
            book.release_resources()
        session = UploadSession(workspace_id=workspace_id, original_filename=filename, file_type=file_type,
                                byte_size=size, sha256=digest.hexdigest(), staging_key=key,
                                sheet_names=sheets, status="uploaded",
                                expires_at=datetime.now(timezone.utc) + timedelta(hours=24))
        db.add(session)
        db.commit()
        db.refresh(session)
        return session
    except Exception:
        path.unlink(missing_ok=True)
        raise


def owned_upload(db: Session, upload_id: int, user_id: int) -> UploadSession:
    session = db.query(UploadSession).join(Workspace).filter(UploadSession.id == upload_id, Workspace.owner_id == user_id).first()
    if session and session.status == "uploaded" and session.expires_at <= datetime.now(timezone.utc):
        if session.staging_key:
            staging_path(session.staging_key).unlink(missing_ok=True)
        session.staging_key = None
        session.status = "expired"
        db.commit()
    if session is None or session.status != "uploaded" or not session.staging_key:
        raise HTTPException(404, "Upload not found or expired")
    return session


def expire_workspace_uploads(db: Session, workspace_id: int) -> None:
    now = datetime.now(timezone.utc)
    expired = db.query(UploadSession).filter(UploadSession.workspace_id == workspace_id,
                      UploadSession.status == "uploaded", UploadSession.expires_at <= now).all()
    for session in expired:
        if session.staging_key:
            staging_path(session.staging_key).unlink(missing_ok=True)
        session.staging_key = None
        session.status = "expired"
    if expired:
        db.commit()


def read_records(session: UploadSession, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    path = staging_path(session.staging_key or "")
    settings = get_settings()
    try:
        if session.file_type in {"csv", "json"}:
            content = path.read_text(encoding="utf-8-sig")
            if session.file_type == "csv":
                reader = csv.DictReader(io.StringIO(content, newline=""))
                headers = list(reader.fieldnames or [])
                rows = list(reader)
                if any(None in row for row in rows):
                    raise ValueError
            else:
                parsed = json.loads(content)
                if not isinstance(parsed, list) or any(not isinstance(item, dict) for item in parsed):
                    raise ValueError
                if any(not isinstance(value, (str, int, float, bool, type(None))) for row in parsed for value in row.values()):
                    raise ValueError
                if any(isinstance(value, float) and not math.isfinite(value) for row in parsed for value in row.values()):
                    raise ValueError
                headers = list(dict.fromkeys(key for row in parsed for key in row.keys()))
                rows = [{key: row.get(key) for key in headers} for row in parsed]
        elif session.file_type == "xlsx":
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if sheet_name not in book.sheetnames:
                book.close()
                raise HTTPException(422, "Select a valid worksheet")
            sheet = book[sheet_name]
            values = sheet.iter_rows(values_only=True)
            first = next(values, None)
            headers = [str(value).strip() if value is not None else "" for value in (first or ())]
            if len(headers) > settings.upload_max_columns:
                book.close()
                raise HTTPException(413, "File exceeds the configured column limit")
            rows = []
            for row in values:
                if len(rows) >= settings.upload_max_rows:
                    book.close()
                    raise HTTPException(413, "File exceeds the configured row limit")
                rows.append(dict(zip(headers, row)))
            book.close()
        else:
            book = xlrd.open_workbook(path, on_demand=True)
            if sheet_name not in book.sheet_names():
                book.release_resources()
                raise HTTPException(422, "Select a valid worksheet")
            sheet = book.sheet_by_name(sheet_name)
            headers = [str(value).strip() for value in (sheet.row_values(0) if sheet.nrows else [])]
            if len(headers) > settings.upload_max_columns or sheet.nrows - 1 > settings.upload_max_rows:
                book.release_resources()
                raise HTTPException(413, "File exceeds the configured row or column limit")
            rows = []
            for row_index in range(1, sheet.nrows):
                values = []
                for column_index in range(len(headers)):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    values.append(value)
                rows.append(dict(zip(headers, values)))
            book.release_resources()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "File could not be parsed; check its encoding and structure") from None
    normalized_headers = [str(name).strip() for name in headers]
    if normalized_headers != headers:
        rows = [{normalized: row.get(original) for original, normalized in zip(headers, normalized_headers)} for row in rows]
        headers = normalized_headers
    if not headers or len(headers) > settings.upload_max_columns or any(not name for name in headers) or len({x.casefold() for x in headers}) != len(headers):
        raise HTTPException(422, "File must contain unique, non-empty column names within the configured limit")
    if len(rows) > settings.upload_max_rows:
        raise HTTPException(413, "File exceeds the configured row limit")
    return headers, rows


def infer_type(values: list[Any]) -> str:
    values = [value for value in values if value not in (None, "")]
    if not values:
        return "text"
    if all(isinstance(value, datetime) for value in values):
        return "datetime"
    if all(isinstance(value, date) and not isinstance(value, datetime) for value in values):
        return "date"
    if all(isinstance(value, bool) for value in values):
        return "boolean"

    def is_integer(value: Any) -> bool:
        if isinstance(value, bool):
            return False
        if isinstance(value, int):
            return True
        if isinstance(value, float):
            return value.is_integer()
        return isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value.strip()) is not None

    if all(is_integer(value) for value in values):
        return "integer"

    def is_decimal(value: Any) -> bool:
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            return False
        try:
            return Decimal(str(value)).is_finite()
        except (InvalidOperation, ValueError, TypeError):
            return False

    if all(is_decimal(value) for value in values):
        return "decimal"
    if all(isinstance(value, str) and value.casefold() in {"true", "false", "yes", "no"} for value in values):
        return "boolean"

    def is_valid_date(value: str) -> bool:
        try:
            date.fromisoformat(value)
            return True
        except ValueError:
            return False

    if all(
        isinstance(value, str)
        and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) is not None
        and is_valid_date(value)
        for value in values
    ):
        return "date"

    def is_valid_datetime(value: str) -> bool:
        try:
            datetime.fromisoformat(value.replace("Z", "+00:00"))
            return True
        except ValueError:
            return False

    if all(
        isinstance(value, str)
        and re.match(r"^\d{4}-\d{2}-\d{2}[T ]", value) is not None
        and is_valid_datetime(value)
        for value in values
    ):
        return "datetime"
    return "text"


def upload_preview(session: UploadSession, sheet_name: str | None) -> dict[str, Any]:
    if session.file_type in {"xlsx", "xls"} and not sheet_name:
        raise HTTPException(422, "Select a worksheet to preview")
    headers, rows = read_records(session, sheet_name)
    columns = [{"name": name, "type": infer_type([row.get(name) for row in rows[:1000]]), "nullable": True} for name in headers]
    return {"upload_id": session.id, "sheet_name": sheet_name, "columns": columns,
            "rows": [{k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in row.items()} for row in rows[:get_settings().upload_preview_rows]],
            "total_rows": len(rows)}


def insert_rows(db: Session, dataset: Dataset, workspace_id: int, rows: list[dict[str, Any]]) -> None:
    table = managed_table(dataset)
    values = validate_rows(dataset, rows)
    for item in values:
        item["workspace_id"] = workspace_id
    if values:
        db.execute(insert(table), values)
        dataset.row_count += len(values)


def create_manual_dataset(db: Session, workspace_id: int, user_id: int, name: str,
                          schema: list[dict[str, Any]], rows: list[dict[str, Any]]) -> Dataset:
    schema = make_schema(schema)
    dataset = Dataset(workspace_id=workspace_id, name=name, source_type="manual", status="ready",
                      schema_json=schema, row_count=0, storage_table_key=f"md_{uuid.uuid4().hex}", source_ref={})
    db.add(dataset)
    db.flush()
    create_storage(db, dataset)
    insert_rows(db, dataset, workspace_id, rows)
    record_revision(db, dataset)
    if rows:
        db.add(IngestionRun(workspace_id=workspace_id, dataset=dataset, actor_user_id=user_id,
                            operation="manual_create", status="completed", rows_seen=len(rows),
                            rows_written=len(rows), error_report=[]))
    db.commit()
    db.refresh(dataset)
    return dataset


def import_upload(db: Session, session: UploadSession, user_id: int, name: str | None, dataset_id: int | None,
                  sheet_name: str | None, append: bool, mapping: dict[str, dict[str, Any]] | None) -> Dataset:
    headers, rows = read_records(session, sheet_name)
    if mapping is None:
        schema = [{"name": header, "type": infer_type([row.get(header) for row in rows[:1000]]), "nullable": True} for header in headers]
        source_to_target = {header: header for header in headers}
    else:
        if set(mapping) != set(headers):
            raise HTTPException(422, "Mapping must specify every source column exactly once")
        schema = [mapping[header] for header in headers]
        source_to_target = {header: mapping[header]["name"].strip() for header in headers}
        if len({name.casefold() for name in source_to_target.values()}) != len(source_to_target):
            raise HTTPException(422, "Mapped column names must be unique")
    rows = [{source_to_target[key]: value for key, value in row.items()} for row in rows]
    if dataset_id is None:
        if not name:
            raise HTTPException(422, "A name is required for a new dataset")
        schema = make_schema(schema)
        dataset = Dataset(workspace_id=session.workspace_id, name=name, source_type="upload", status="ready",
                          schema_json=schema, row_count=0, storage_table_key=f"md_{uuid.uuid4().hex}",
                          source_ref={"upload_id": session.id, "filename": session.original_filename, "sha256": session.sha256})
        db.add(dataset)
        db.flush()
        create_storage(db, dataset)
    else:
        dataset = db.query(Dataset).join(Workspace).filter(Dataset.id == dataset_id,
                    Dataset.workspace_id == session.workspace_id, Workspace.owner_id == user_id).with_for_update().first()
        if not dataset:
            raise HTTPException(404, "Dataset not found")
        if dataset.source_type == "database" or not append:
            raise HTTPException(409, "This dataset cannot be appended to")
        if [{k: c[k] for k in ("name", "type")} for c in schema] != [{k: c[k] for k in ("name", "type")} for c in dataset.schema_json]:
            raise HTTPException(422, "Imported columns must match the dataset schema in order")
    try:
        insert_rows(db, dataset, session.workspace_id, rows)
        record_revision(db, dataset)
        db.add(IngestionRun(workspace_id=session.workspace_id, dataset=dataset, upload_id=session.id,
                            actor_user_id=user_id, operation="append" if dataset_id else "import", status="completed",
                            rows_seen=len(rows), rows_written=len(rows), error_report=[]))
        session.status = "imported"
        key = session.staging_key
        session.staging_key = None
        db.commit()
        if key:
            staging_path(key).unlink(missing_ok=True)
        db.refresh(dataset)
        return dataset
    except Exception:
        db.rollback()
        raise


def read_dataset_rows(db: Session, dataset: Dataset, workspace_id: int, limit: int, offset: int) -> list[dict[str, Any]]:
    table = managed_table(dataset)
    rows = db.execute(select(table).where(table.c.workspace_id == workspace_id).order_by(table.c.row_id).limit(limit).offset(offset)).mappings().all()
    reverse = {spec["key"]: spec["name"] for spec in dataset.schema_json}
    return [{"id": row["row_id"], **{reverse[key]: row[key] for key in reverse}} for row in rows]


def update_dataset_row(db: Session, dataset: Dataset, workspace_id: int, row_id: str, values: dict[str, Any]) -> None:
    table = managed_table(dataset)
    specs = {spec["name"]: spec for spec in dataset.schema_json}
    if not values:
        raise HTTPException(422, "values must contain at least one column")
    unknown_columns = set(values) - set(specs)
    if unknown_columns:
        names = sorted(str(name)[:64] for name in unknown_columns)[:10]
        raise HTTPException(422, f"Unknown columns: {', '.join(names)}")
    mapped = {}
    for name, value in values.items():
        try: mapped[specs[name]["key"]] = coerce_value(value, specs[name]["type"], specs[name]["nullable"])
        except (ValueError, TypeError, InvalidOperation): raise HTTPException(422, f"Invalid value for column {name}") from None
    result = db.execute(update(table).where(table.c.row_id == row_id, table.c.workspace_id == workspace_id).values(**mapped))
    if not result.rowcount: raise HTTPException(404, "Dataset row not found")
    record_revision(db, dataset)


def delete_dataset_row(db: Session, dataset: Dataset, workspace_id: int, row_id: str) -> None:
    table = managed_table(dataset)
    result = db.execute(delete(table).where(table.c.row_id == row_id, table.c.workspace_id == workspace_id))
    if not result.rowcount: raise HTTPException(404, "Dataset row not found")
    dataset.row_count = max(0, dataset.row_count - 1)
    record_revision(db, dataset)

