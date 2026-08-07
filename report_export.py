import io
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fpdf import FPDF
import arabic_reshaper
from bidi.algorithm import get_display

# Path to the bundled Arabic-capable font used for PDF export
FONT_DIR = os.path.join(os.path.dirname(__file__), "fonts")
ARABIC_FONT_PATH = os.path.join(FONT_DIR, "Amiri-Regular.ttf")


# Splits the LLM's markdown-style report text into simple typed lines
# (heading / bullet / paragraph) so both exporters can render it consistently
def parse_report_lines(report_text: str) -> list:
    lines = []
    for raw_line in report_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            lines.append({"type": "heading", "text": line[3:].strip()})
        elif line.startswith("# "):
            lines.append({"type": "heading", "text": line[2:].strip()})
        elif line.startswith(("- ", "* ")):
            lines.append({"type": "bullet", "text": line[2:].strip()})
        else:
            lines.append({"type": "paragraph", "text": line})
    return lines


# ---------- Word (.docx) export ----------
# Builds a Word document from the report text, applying right-to-left
# paragraph direction automatically when the report is in Arabic
def generate_docx(report_text: str, language: str) -> bytes:
    doc = Document()
    is_arabic = language == "Arabic"

    for item in parse_report_lines(report_text):
        if item["type"] == "heading":
            p = doc.add_heading(item["text"], level=2)
        elif item["type"] == "bullet":
            p = doc.add_paragraph(item["text"], style="List Bullet")
        else:
            p = doc.add_paragraph(item["text"])

        if is_arabic:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            # Mark the paragraph as right-to-left so Word displays Arabic correctly
            p_format = p.paragraph_format
            p_format.element.get_or_add_pPr().set(qn("w:bidi"), "1")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------- PDF export ----------
# Builds a PDF from the report text. English uses the built-in font;
# Arabic requires shaping (arabic_reshaper) and right-to-left reordering
# (python-bidi) plus a Unicode font, since PDF has no native RTL support.
def generate_pdf(report_text: str, language: str) -> bytes:
    is_arabic = language == "Arabic"
    pdf = FPDF()
    pdf.add_page()

    if is_arabic:
        if not os.path.exists(ARABIC_FONT_PATH):
            raise FileNotFoundError(
                "Arabic font not found. Make sure 'fonts/Amiri-Regular.ttf' "
                "is present next to your project files."
            )
        pdf.add_font("Amiri", "", ARABIC_FONT_PATH)
        pdf.set_font("Amiri", size=13)
    else:
        pdf.set_font("Helvetica", size=12)

    for item in parse_report_lines(report_text):
        text = item["text"]

        if is_arabic:
            # Reshape Arabic letters into their joined forms, then reorder
            # for right-to-left display before handing text to fpdf
            reshaped = arabic_reshaper.reshape(text)
            text = get_display(reshaped)

        if item["type"] == "heading":
            pdf.set_font(pdf.font_family, "B" if not is_arabic else "", 15)
            pdf.multi_cell(0, 10, text, align="R" if is_arabic else "L")
            pdf.set_font(pdf.font_family, "", 13 if is_arabic else 12)
        elif item["type"] == "bullet":
            bullet_text = f"- {text}" if not is_arabic else f"{text} -"
            pdf.multi_cell(0, 8, bullet_text, align="R" if is_arabic else "L")
        else:
            pdf.multi_cell(0, 8, text, align="R" if is_arabic else "L")
        pdf.ln(1)

    return bytes(pdf.output())
